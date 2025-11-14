"""
Excel Handler Module
Handles all Excel file operations including reading invoice numbers and updating status.
"""

import openpyxl
from openpyxl import load_workbook, Workbook
import logging
from datetime import datetime


class ExcelHandler:
    """
    Manages Excel file operations for invoice checking.
    Reads invoice numbers and updates their status.
    """
    
    def __init__(self, file_path):
        """
        Initialize the Excel handler with a file path.
        
        Args:
            file_path (str): Path to the Excel file
        """
        self.file_path = file_path
        self.workbook = None
        self.worksheet = None
        self.column_indices = {}  # Dictionary to store column indices
        self.status_column = None
        self.timestamp_column = None
        self.value_of_purchases_column = None
        
    def load_excel(self):
        """
        Load the Excel file and validate required columns.
        
        Returns:
            bool: True if loaded successfully, False otherwise
        """
        try:
            self.workbook = load_workbook(self.file_path)
            self.worksheet = self.workbook.active
            
            # Find all available columns (case-insensitive)
            headers = [cell.value for cell in self.worksheet[1]]
            headers_lower = [h.lower() if h else None for h in headers]
            
            # Expected columns mapping
            column_names = [
                'seller registration no',
                'seller registration no.',
                'seller name',
                'number',
                'date',
                'source authority',
                'sr.no',
                'purchase type',
                'rate',
                'value of purchases',
                'sales tax/ fed in st mode'
            ]
            
            # Find columns in the header (case-insensitive)
            for col_name in column_names:
                if col_name in headers_lower:
                    idx = headers_lower.index(col_name)
                    self.column_indices[col_name] = idx + 1  # 1-indexed
                    logging.info(f"Found column '{headers[idx]}' at column {idx + 1}")
            
            # Check if at least the main column exists
            seller_reg_found = any(k in self.column_indices for k in ['seller registration no', 'seller registration no.'])
            
            if not seller_reg_found:
                logging.error("Column 'Seller Registration No' not found in Excel file")
                logging.error(f"Available columns: {headers}")
                return False
            
            # Add Status column if it doesn't exist
            if 'Status' not in headers:
                status_col_idx = len(headers) + 1
                self.worksheet.cell(row=1, column=status_col_idx, value='Status')
                self.status_column = status_col_idx
                logging.info("Added 'Status' column to Excel file")
            else:
                self.status_column = headers.index('Status') + 1
            
            # Add Checked_On column if it doesn't exist
            if 'Checked_On' not in headers:
                timestamp_col_idx = len(headers) + (2 if 'Status' not in headers else 1)
                self.worksheet.cell(row=1, column=timestamp_col_idx, value='Checked_On')
                self.timestamp_column = timestamp_col_idx
                logging.info("Added 'Checked_On' column to Excel file")
            else:
                headers = [cell.value for cell in self.worksheet[1]]  # Refresh headers
                self.timestamp_column = headers.index('Checked_On') + 1
            
            # Add Value_of_Purchases column if it doesn't exist
            headers = [cell.value for cell in self.worksheet[1]]  # Refresh headers
            if 'Value_of_Purchases' not in headers:
                value_col_idx = len(headers) + 1
                self.worksheet.cell(row=1, column=value_col_idx, value='Value_of_Purchases')
                self.value_of_purchases_column = value_col_idx
                logging.info("Added 'Value_of_Purchases' column to Excel file")
            else:
                self.value_of_purchases_column = headers.index('Value_of_Purchases') + 1
            
            self.workbook.save(self.file_path)
            logging.info(f"Excel file loaded successfully: {self.file_path}")
            logging.info(f"Column mapping: {self.column_indices}")
            return True
            
        except Exception as e:
            logging.error(f"Error loading Excel file: {str(e)}")
            return False
    
    def get_invoice_numbers(self):
        """
        Retrieve all invoice data from the Excel file.
        Returns a list of dictionaries with all relevant data for each row.
        
        Returns:
            list: List of dictionaries containing invoice data
        """
        invoices = []
        
        try:
            # Start from row 2 (skip header)
            for row in range(2, self.worksheet.max_row + 1):
                # Get Seller Registration No (required)
                seller_reg_col = self.column_indices.get('seller registration no') or self.column_indices.get('seller registration no.')
                
                if seller_reg_col:
                    seller_registration_no = self.worksheet.cell(row=row, column=seller_reg_col).value
                    
                    if seller_registration_no:  # Skip empty rows
                        # Build invoice data dictionary
                        invoice_data = {
                            'row': row,
                            'seller_registration_no': str(seller_registration_no).strip()
                        }
                        
                        # Get optional columns if they exist
                        if 'seller name' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['seller name']).value
                            invoice_data['seller_name'] = str(val).strip() if val else 'N/A'
                        
                        if 'number' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['number']).value
                            invoice_data['number'] = str(val).strip() if val else 'N/A'
                        
                        if 'date' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['date']).value
                            invoice_data['date'] = str(val).strip() if val else 'N/A'
                        
                        if 'source authority' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['source authority']).value
                            invoice_data['source_authority'] = str(val).strip() if val else 'FBR'
                        
                        if 'sr.no' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['sr.no']).value
                            invoice_data['sr_no'] = str(val).strip() if val else str(row - 1)
                        
                        if 'sales tax/ fed in st mode' in self.column_indices:
                            val = self.worksheet.cell(row=row, column=self.column_indices['sales tax/ fed in st mode']).value
                            invoice_data['sales_tax_fed_st_mode'] = str(val).strip() if val else 'N/A'
                        
                        invoices.append(invoice_data)
            
            logging.info(f"Retrieved {len(invoices)} invoice records from Excel")
            return invoices
            
        except Exception as e:
            logging.error(f"Error reading invoice data: {str(e)}")
            return []
    
    def update_invoice_status(self, row_number, status, value_of_purchases=None):
        """
        Update the status of an invoice in the Excel file.
        
        Args:
            row_number (int): Row number in Excel (1-indexed)
            status (str): Status to update (Claimed/Not Claimed/Error)
            value_of_purchases (str, optional): Value of Purchases from FBR portal
        """
        try:
            # Update Status column
            self.worksheet.cell(row=row_number, column=self.status_column, value=status)
            
            # Update timestamp
            timestamp = datetime.now().strftime('%Y-%m-%d %I:%M:%S %p')
            self.worksheet.cell(row=row_number, column=self.timestamp_column, value=timestamp)
            
            # Update Value of Purchases if provided
            if value_of_purchases and self.value_of_purchases_column:
                self.worksheet.cell(row=row_number, column=self.value_of_purchases_column, value=value_of_purchases)
                logging.info(f"Updated row {row_number} with Value of Purchases: {value_of_purchases}")
            
            # Save immediately to prevent data loss
            self.workbook.save(self.file_path)
            
            logging.info(f"Updated row {row_number} with status: {status}")
            
        except Exception as e:
            logging.error(f"Error updating invoice status: {str(e)}")
    
    def close(self):
        """
        Close the Excel workbook.
        """
        try:
            if self.workbook:
                self.workbook.close()
                logging.info("Excel workbook closed")
        except Exception as e:
            logging.error(f"Error closing workbook: {str(e)}")
