"""
Sample Excel File Generator
Creates a sample invoices.xlsx file with test invoice numbers.
"""

from openpyxl import Workbook
import os


def create_sample_excel():
    """
    Create a sample Excel file with invoice numbers for testing.
    """
    # Create a new workbook
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"
    
    # Add headers
    ws['A1'] = 'InvoiceNumber'
    
    # Sample invoice numbers (these are dummy numbers for testing)
    sample_invoices = [
        '1234567890123',
        '2345678901234',
        '3456789012345',
        '4567890123456',
        '5678901234567',
        '6789012345678',
        '7890123456789',
        '8901234567890',
        '9012345678901',
        '0123456789012',
    ]
    
    # Add invoice numbers to Excel
    for idx, invoice in enumerate(sample_invoices, start=2):
        ws[f'A{idx}'] = invoice
    
    # Save the file
    file_path = os.path.join(os.path.dirname(__file__), 'invoices.xlsx')
    wb.save(file_path)
    
    print(f"✅ Sample Excel file created: {file_path}")
    print(f"📋 Added {len(sample_invoices)} sample invoice numbers")
    print("\nYou can now:")
    print("1. Open invoices.xlsx and replace with your actual invoice numbers")
    print("2. Or run: python main.py to start the application")


if __name__ == "__main__":
    create_sample_excel()
